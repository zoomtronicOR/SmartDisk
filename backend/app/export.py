"""Izvoz tabela u CSV, XLSX i PDF (spec. §10.5, §10.7)."""

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

FORMATS = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

# DejaVu ima sva slova srpske latinice (č, ć, đ, š, ž); ugrađeni PDF fontovi nemaju đ i ć.
_FONT_DIRS = [Path("/usr/share/fonts/truetype/dejavu"), Path("C:/Windows/Fonts")]


@dataclass
class Column:
    key: str
    label: str
    kind: str = "text"  # text / int / float / datetime / bytes / percent


@dataclass
class Report:
    title: str
    columns: list[Column]
    rows: list[dict[str, Any]]
    subtitle: str | None = None
    notes: list[str] = field(default_factory=list)


def _cell(value: Any, column: Column, tz: ZoneInfo, for_excel: bool = False) -> Any:
    if value is None:
        return "" if not for_excel else None
    if isinstance(value, datetime):
        local = value.astimezone(tz)
        return local.replace(tzinfo=None) if for_excel else local.strftime("%d.%m.%Y. %H:%M")
    if isinstance(value, date):
        return value if for_excel else value.strftime("%d.%m.%Y.")
    if isinstance(value, bool):
        return "da" if value else "ne"
    if isinstance(value, (list, tuple)):
        return "; ".join(str(v) for v in value)
    if isinstance(value, float) and not for_excel:
        return f"{value:.2f}".replace(".", ",")
    return value


def to_csv(report: Report, tz: ZoneInfo) -> bytes:
    """CSV sa `;` kao separatorom i BOM-om, da ga Excel sa srpskim podešavanjima otvori ispravno."""
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n")
    writer.writerow([c.label for c in report.columns])
    for row in report.rows:
        writer.writerow([_cell(row.get(c.key), c, tz) for c in report.columns])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_xlsx(report: Report, tz: ZoneInfo) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31].replace("/", "-")
    ws.append([report.title])
    ws["A1"].font = Font(bold=True, size=13)
    if report.subtitle:
        ws.append([report.subtitle])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([c.label for c in report.columns])
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for row in report.rows:
        ws.append([_cell(row.get(c.key), c, tz, for_excel=True) for c in report.columns])
    for i, column in enumerate(report.columns, start=1):
        letter = get_column_letter(i)
        if column.kind == "datetime":
            for cell in ws[letter][header_row:]:
                cell.number_format = "dd.mm.yyyy hh:mm"
        width = max([len(column.label)] + [len(str(_cell(r.get(column.key), column, tz))) for r in report.rows[:500]])
        ws.column_dimensions[letter].width = min(60, max(8, width + 2))
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for note in report.notes:
        ws.append([])
        ws.append([note])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _register_fonts() -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if "DejaVuSans" in pdfmetrics.getRegisteredFontNames():
        return "DejaVuSans", "DejaVuSans-Bold"
    for directory in _FONT_DIRS:
        regular, bold = directory / "DejaVuSans.ttf", directory / "DejaVuSans-Bold.ttf"
        if regular.exists() and bold.exists():
            pdfmetrics.registerFont(TTFont("DejaVuSans", str(regular)))
            pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(bold)))
            return "DejaVuSans", "DejaVuSans-Bold"
    return "Helvetica", "Helvetica-Bold"


def to_pdf(report: Report, tz: ZoneInfo, generated_by: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    font, bold = _register_fonts()
    wide = len(report.columns) > 6
    pagesize = landscape(A4) if wide else A4
    out = io.BytesIO()
    generated = datetime.now(tz).strftime("%d.%m.%Y. %H:%M")

    def footer(canvas, doc):  # noqa: ANN001
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.HexColor("#52514e"))
        canvas.drawString(12 * mm, 8 * mm, f"SmartDisk · {report.title} · generisano {generated}"
                          + (f" · {generated_by}" if generated_by else ""))
        canvas.drawRightString(pagesize[0] - 12 * mm, 8 * mm, f"strana {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=pagesize, leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=report.title, author="SmartDisk")
    title = ParagraphStyle("t", fontName=bold, fontSize=14, leading=18, spaceAfter=2 * mm)
    sub = ParagraphStyle("s", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor("#52514e"))
    cell = ParagraphStyle("c", fontName=font, fontSize=7, leading=8.5)
    head = ParagraphStyle("h", fontName=bold, fontSize=7, leading=8.5)

    story: list[Any] = [Paragraph(_esc(report.title), title)]
    if report.subtitle:
        story.append(Paragraph(_esc(report.subtitle), sub))
    story.append(Spacer(1, 4 * mm))
    if not report.rows:
        story.append(Paragraph("Nema podataka za izabrani period.", sub))
    else:
        data = [[Paragraph(_esc(c.label), head) for c in report.columns]]
        for row in report.rows:
            data.append([Paragraph(_esc(str(_cell(row.get(c.key), c, tz))), cell) for c in report.columns])
        usable = pagesize[0] - 24 * mm
        weights = [3 if c.kind == "text" else 2 for c in report.columns]
        widths = [usable * w / sum(weights) for w in weights]
        table = LongTable(data, colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0efec")),
            ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#c3c2b7")),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#e1e0d9")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]))
        story.append(table)
    for note in report.notes:
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(_esc(note), sub))
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def _esc(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def render(report: Report, fmt: str, tz: ZoneInfo, generated_by: str | None = None) -> bytes:
    if fmt == "csv":
        return to_csv(report, tz)
    if fmt == "xlsx":
        return to_xlsx(report, tz)
    if fmt == "pdf":
        return to_pdf(report, tz, generated_by)
    raise ValueError(fmt)
