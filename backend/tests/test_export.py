"""Izvoz: srpska slova moraju preživeti CSV, XLSX i PDF."""

import io
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from app.export import Column, Report, render

TZ = ZoneInfo("Europe/Belgrade")
REPORT = Report(
    "Izveštaj – čćžšđ ČĆŽŠĐ",
    [Column("name", "Uređaj"), Column("t", "Očitano", "datetime"), Column("v", "Vrednost", "float"),
     Column("ok", "Prisutan")],
    [{"name": "Proxmox i5 – đak", "t": datetime(2026, 10, 28, 20, 0, tzinfo=timezone.utc), "v": 1.5, "ok": True},
     {"name": "TrueNAS", "t": None, "v": None, "ok": False}],
    "Period: test",
    ["Napomena: prazno polje znači „nije dostupno”."],
)


def test_csv():
    data = render(REPORT, "csv", TZ)
    text = data.decode("utf-8")
    assert text.startswith("﻿") and "Uređaj;Očitano;Vrednost;Prisutan" in text
    assert "Proxmox i5 – đak;28.10.2026. 21:00;1,50;da" in text  # lokalno vreme, decimalni zarez
    assert "TrueNAS;;;ne" in text


def test_xlsx():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(render(REPORT, "xlsx", TZ)))
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Izveštaj – čćžšđ ČĆŽŠĐ" in values and "Proxmox i5 – đak" in values
    assert datetime(2026, 10, 28, 21, 0) in values


def test_pdf():
    data = render(REPORT, "pdf", TZ, "admin")
    assert data.startswith(b"%PDF") and len(data) > 1000
